"""
report.py — Generates a professional Word (.docx) report from research JSON.
Uses python-docx. Mirrors the structure of the 4Imprint report we built earlier.
"""

import os
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ─── Colour palette ────────────────────────────────────────────────────────
BRAND_BLUE  = RGBColor(0x1F, 0x4E, 0x79)
MID_BLUE    = RGBColor(0x2E, 0x75, 0xB6)
LIGHT_BLUE  = RGBColor(0xD6, 0xE4, 0xF0)
GREEN_BG    = RGBColor(0xE2, 0xEF, 0xDA)
AMBER_BG    = RGBColor(0xFF, 0xF2, 0xCC)
DARK_GREY   = RGBColor(0x40, 0x40, 0x40)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GREY  = RGBColor(0xF2, 0xF2, 0xF2)


def _hex(rgb: RGBColor) -> str:
    return str(rgb)  # RGBColor.__str__ returns the 6-char uppercase hex string


def _set_cell_bg(cell, hex_color: str):
    """Set table cell background colour."""
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color)
    tcPr.append(shd)


def _set_cell_borders(cell, color: str = "CCCCCC"):
    """Set single-line borders on a table cell."""
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"),   "single")
        el.set(qn("w:sz"),    "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def _cell_margins(cell, top=80, bottom=80, left=120, right=120):
    """Set internal cell padding (in DXA/twentieths of a point)."""
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    mar  = OxmlElement("w:tcMar")
    for side, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"),    str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def _para_border_bottom(para, color: str = "2E75B6", sz: int = 6):
    """Add a bottom border to a paragraph (used for section headings)."""
    pPr   = para._p.get_or_add_pPr()
    pBdr  = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"),   "single")
    bottom.set(qn("w:sz"),    str(sz))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def _set_col_widths(table, widths_twips: list[int]):
    """Set explicit column widths on a table (in twips)."""
    tbl   = table._tbl
    tblPr = tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"),    str(sum(widths_twips)))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tc   = cell._tc
            tcPr = tc.get_or_add_tcPr()
            tcW  = OxmlElement("w:tcW")
            tcW.set(qn("w:w"),    str(widths_twips[i]))
            tcW.set(qn("w:type"), "dxa")
            tcPr.append(tcW)


# ─── Document helpers ──────────────────────────────────────────────────────
def _add_heading1(doc: Document, text: str):
    para = doc.add_paragraph()
    para.style = doc.styles["Heading 1"]
    run = para.add_run(text)
    run.font.color.rgb = BRAND_BLUE
    run.font.size = Pt(16)
    run.bold = True
    _para_border_bottom(para)
    para.paragraph_format.space_before = Pt(18)
    para.paragraph_format.space_after  = Pt(9)
    return para


def _add_heading2(doc: Document, text: str):
    para = doc.add_paragraph()
    para.style = doc.styles["Heading 2"]
    run = para.add_run(text)
    run.font.color.rgb = MID_BLUE
    run.font.size = Pt(13)
    run.bold = True
    para.paragraph_format.space_before = Pt(14)
    para.paragraph_format.space_after  = Pt(6)
    return para


def _add_heading3(doc: Document, text: str):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.color.rgb = DARK_GREY
    run.font.size = Pt(11)
    run.bold = True
    para.paragraph_format.space_before = Pt(10)
    para.paragraph_format.space_after  = Pt(4)
    return para


def _add_body(doc: Document, text: str, italic: bool = False):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.size = Pt(11)
    run.font.color.rgb = DARK_GREY
    run.italic = italic
    para.paragraph_format.space_after = Pt(6)
    return para


def _add_bullet(doc: Document, text: str):
    para = doc.add_paragraph(style="List Bullet")
    run = para.add_run(text)
    run.font.size = Pt(11)
    run.font.color.rgb = DARK_GREY
    para.paragraph_format.space_after = Pt(4)
    return para


def _score_color(score: float) -> RGBColor:
    if score >= 7.5:
        return GREEN_BG
    elif score >= 5.5:
        return AMBER_BG
    else:
        return RGBColor(0xFC, 0xE4, 0xD6)  # red-ish


def _add_score_box(doc: Document, label: str, score: float, max_score: int = 10):
    """Add a single-row coloured score summary box."""
    color = _score_color(score)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    _set_col_widths(table, [7200, 2160])

    left, right = table.rows[0].cells
    _set_cell_bg(left,  _hex(color))
    _set_cell_bg(right, _hex(color))
    _set_cell_borders(left)
    _set_cell_borders(right)
    _cell_margins(left);  _cell_margins(right)

    lp = left.paragraphs[0]
    lr = lp.add_run(label)
    lr.font.size = Pt(11); lr.bold = True; lr.font.color.rgb = DARK_GREY

    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = rp.add_run(f"{score} / {max_score}")
    rr.font.size = Pt(13); rr.bold = True; rr.font.color.rgb = BRAND_BLUE

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def _add_scorecard_table(doc: Document, tenet_scores: dict, overall_score: float):
    """Add the main 4-row scorecard table."""
    headers = ["Tenet", "Score", "Weight", "Weighted"]
    rows = [
        ("Durability (Moat Strength)",        tenet_scores["durability"]["score"], "40%",
         round(tenet_scores["durability"]["score"] * 0.40, 1)),
        ("Compounding Capability",            tenet_scores["compounding"]["score"], "35%",
         round(tenet_scores["compounding"]["score"] * 0.35, 1)),
        ("Fragility (10 = least fragile)",    tenet_scores["fragility"]["score"],  "25%",
         round(tenet_scores["fragility"]["score"] * 0.25, 1)),
        ("OVERALL QUALITY SCORE",             overall_score,                        "100%",
         overall_score),
    ]
    col_widths = [4500, 1500, 1620, 1740]
    table = doc.add_table(rows=1 + len(rows), cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    _set_col_widths(table, col_widths)

    # Header row
    hrow = table.rows[0]
    for i, h in enumerate(headers):
        cell = hrow.cells[i]
        _set_cell_bg(cell, _hex(BRAND_BLUE))
        _set_cell_borders(cell, _hex(MID_BLUE))
        _cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.size = Pt(10); r.bold = True; r.font.color.rgb = WHITE

    # Data rows
    for ri, (label, score, weight, weighted) in enumerate(rows):
        row = table.rows[ri + 1]
        is_total = label.startswith("OVERALL")
        fill = _hex(LIGHT_BLUE) if is_total else (
            _hex(GREEN_BG) if score >= 7.5 else
            _hex(AMBER_BG) if score >= 5.5 else "FCE4D6"
        )
        values = [label, str(score), weight, str(weighted)]
        aligns = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER,
                  WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER]
        for ci, (val, align) in enumerate(zip(values, aligns)):
            cell = row.cells[ci]
            _set_cell_bg(cell, fill)
            _set_cell_borders(cell)
            _cell_margins(cell)
            p = cell.paragraphs[0]
            p.alignment = align
            r = p.add_run(val)
            r.font.size = Pt(10)
            r.bold = is_total or (ci == 0)
            r.font.color.rgb = BRAND_BLUE if ci in (1, 3) and is_total else DARK_GREY

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table


# ─── Main entry point ──────────────────────────────────────────────────────
def generate_report(data: dict, output_path: str) -> str:
    """
    Generate a Word report from research data dict.
    Returns the path of the written .docx file.
    """
    doc = Document()

    # ── Page setup: US Letter, 1-inch margins ─────────────────────────────
    for section in doc.sections:
        section.page_width  = Twips(12240)
        section.page_height = Twips(15840)
        section.left_margin   = Inches(1)
        section.right_margin  = Inches(1)
        section.top_margin    = Inches(1)
        section.bottom_margin = Inches(1)

    # ── Default font ──────────────────────────────────────────────────────
    doc.styles["Normal"].font.name  = "Arial"
    doc.styles["Normal"].font.size  = Pt(11)

    company    = data.get("company_name", "Unknown Company")
    score      = data.get("overall_score", 0)
    score_label = data.get("overall_label", "")
    date_str   = data.get("research_date", datetime.today().strftime("%Y-%m-%d"))
    tscores    = data.get("tenet_scores", {})
    questions  = data.get("questions", [])

    # ════════════════════════════════════════════════════════════════════════
    # COVER PAGE
    # ════════════════════════════════════════════════════════════════════════
    doc.add_paragraph().paragraph_format.space_before = Pt(60)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run(company.upper())
    tr.font.size = Pt(28); tr.bold = True; tr.font.color.rgb = BRAND_BLUE
    title.paragraph_format.space_after = Pt(12)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("Investment Quality Assessment")
    sr.font.size = Pt(16); sr.font.color.rgb = MID_BLUE
    sub.paragraph_format.space_after = Pt(36)

    doc.add_paragraph().paragraph_format.space_before = Pt(24)
    score_para = doc.add_paragraph()
    score_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    score_para.add_run("Overall Quality Score\n").font.size = Pt(12)
    big = score_para.add_run(f"{score} / 10")
    big.font.size = Pt(36); big.bold = True; big.font.color.rgb = BRAND_BLUE
    score_para.paragraph_format.space_after = Pt(8)

    label_para = doc.add_paragraph()
    label_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lr = label_para.add_run(score_label)
    lr.font.size = Pt(12); lr.italic = True; lr.font.color.rgb = DARK_GREY
    label_para.paragraph_format.space_after = Pt(48)

    doc.add_paragraph().paragraph_format.space_before = Pt(48)
    date_para = doc.add_paragraph()
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_para.add_run(f"Date: {date_str}").font.size = Pt(11)

    doc.add_page_break()

    # ════════════════════════════════════════════════════════════════════════
    # EXECUTIVE SUMMARY
    # ════════════════════════════════════════════════════════════════════════
    _add_heading1(doc, "Executive Summary")
    _add_body(doc, data.get("executive_summary", ""))

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    _add_heading3(doc, "Key Findings")
    for finding in data.get("key_findings", []):
        _add_bullet(doc, finding)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    _add_heading3(doc, "Scorecard Summary")
    _add_scorecard_table(doc, tscores, score)

    doc.add_page_break()

    # ════════════════════════════════════════════════════════════════════════
    # TENET SECTIONS
    # ════════════════════════════════════════════════════════════════════════
    tenets = [
        ("Durability",            "durability",  [q for q in questions if q["id"].startswith("1.")]),
        ("Compounding Capability","compounding", [q for q in questions if q["id"].startswith("2.")]),
        ("Fragility",             "fragility",   [q for q in questions if q["id"].startswith("3.")]),
    ]

    for tenet_label, tenet_key, tenet_qs in tenets:
        _add_heading1(doc, f"Section: {tenet_label}")

        ts = tscores.get(tenet_key, {})
        if ts.get("rationale"):
            _add_body(doc, ts["rationale"], italic=True)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)

        for q in tenet_qs:
            q_id        = q.get("id", "")
            sub_tenet   = q.get("sub_tenet", q.get("question", ""))
            answer      = q.get("answer", "")
            judgement   = q.get("judgement", "")
            q_score     = q.get("score", 0)

            _add_heading2(doc, f"{q_id} — {sub_tenet}")

            if answer:
                _add_body(doc, answer)

            # Inline score chip
            chip = doc.add_paragraph()
            chip.paragraph_format.space_before = Pt(4)
            chip.paragraph_format.space_after  = Pt(10)
            chip_run = chip.add_run(f"Judgement: {judgement}   |   Score: {q_score} / 10")
            chip_run.font.size = Pt(10)
            chip_run.bold = True
            chip_run.font.color.rgb = MID_BLUE

        # Tenet score box
        ts_score = ts.get("score", 0)
        doc.add_paragraph().paragraph_format.space_before = Pt(6)
        _add_score_box(doc, f"{tenet_label} Score", ts_score)
        doc.add_page_break()

    # ════════════════════════════════════════════════════════════════════════
    # OVERALL VERDICT
    # ════════════════════════════════════════════════════════════════════════
    _add_heading1(doc, "Overall Verdict")

    _add_heading2(doc, "Weighted Scorecard")
    _add_scorecard_table(doc, tscores, score)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    _add_heading2(doc, "Investment Thesis")
    _add_body(doc, data.get("investment_thesis", ""))

    _add_heading2(doc, "Key Risks to Monitor")
    for risk in data.get("key_risks", []):
        _add_bullet(doc, risk)

    _add_heading2(doc, "Conviction")
    conv     = data.get("conviction", "")
    conv_rat = data.get("conviction_rationale", "")
    _add_body(doc, f"Conviction level: {conv}")
    if conv_rat:
        _add_body(doc, conv_rat)

    doc.add_page_break()

    # ════════════════════════════════════════════════════════════════════════
    # SOURCES
    # ════════════════════════════════════════════════════════════════════════
    _add_heading1(doc, "Sources")
    sources = data.get("sources", [])
    if sources:
        col_widths = [3000, 6360]
        table = doc.add_table(rows=1 + len(sources), cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        _set_col_widths(table, col_widths)

        # Header
        hrow = table.rows[0]
        for ci, htext in enumerate(["Source", "URL"]):
            cell = hrow.cells[ci]
            _set_cell_bg(cell, _hex(BRAND_BLUE))
            _set_cell_borders(cell, _hex(MID_BLUE))
            _cell_margins(cell)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(htext)
            r.font.size = Pt(10); r.bold = True; r.font.color.rgb = WHITE

        for ri, src in enumerate(sources):
            row = table.rows[ri + 1]
            fill = _hex(LIGHT_GREY) if ri % 2 == 0 else _hex(WHITE)

            title_cell = row.cells[0]
            _set_cell_bg(title_cell, fill)
            _set_cell_borders(title_cell)
            _cell_margins(title_cell)
            tp = title_cell.paragraphs[0]
            tp.add_run(src.get("title", "")).font.size = Pt(9)

            url_cell = row.cells[1]
            _set_cell_bg(url_cell, fill)
            _set_cell_borders(url_cell)
            _cell_margins(url_cell)
            up = url_cell.paragraphs[0]
            ur = up.add_run(src.get("url", ""))
            ur.font.size = Pt(9); ur.font.color.rgb = MID_BLUE; ur.italic = True
    else:
        _add_body(doc, "No sources recorded.")

    # ── Save ─────────────────────────────────────────────────────────────
    parent = os.path.dirname(output_path)
    if parent:
        os.makedirs(parent, exist_ok=True)
    doc.save(output_path)
    return output_path
